"""Build the short management overview from verified project capabilities."""

from html import escape
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
PURPLE = "4F2D7F"

INTRO = (
    "Learning Space is an internal Learning Tracker that helps managers plan and assign "
    "employee learning, monitor progress, and retain completion evidence in one place. "
    "Employees can view their learning plan, update permitted progress, and upload certificates. "
    "The project covers Phase 1 only."
)
ROLES = [
    ("Manager (also administrator)", "Manages employees, courses, tracks and assignments; reviews team progress, certificates and audit history."),
    ("Employee", "Views own learning, deadlines and history; updates allowed progress and provides completion evidence."),
]
FEATURES = [
    ("Authentication", "Email/password login, session renewal, logout and protected access."),
    ("Employee management", "Create, edit, search and activate/deactivate employees; maintain department and team membership."),
    ("Course management", "Maintain course code, description, category, provider and duration; deactivate while preserving history."),
    ("Learning tracks", "Group courses into ordered learning paths; add, remove and reorder courses."),
    ("Assignments", "Assign a course or track to an employee; set deadlines, update policies and cancel assignments."),
    ("Employee dashboard", "Personal learning plan, status summaries, filters, target dates and learning history."),
    ("Completion tracking", "Track Not Started, Enrolled, In Progress and Completed; validate dates and allow manager corrections."),
    ("Certificates / evidence", "Upload and download PDF, PNG or JPEG evidence with ownership, file type and size checks."),
    ("Manager dashboard", "Team progress, completion rates, pending/overdue learning and capability-building signals."),
    ("Audit logging", "Record who changed employee, course, assignment, progress and certificate information."),
]
TEAM_NOTE = "Team creation is currently available through the manager API; employee forms support selecting an existing team."
WORKFLOW = "Create employee and courses → build a track → assign learning and a target date → update progress → record completion/evidence → review team results."
VALUE = "Managers gain a consolidated view of development needs and overdue learning. Employees have clear learning expectations and an accessible completion history."
SCOPE = "LMS/iLearn imports, CSV/Excel processing, staging and exception reports are outside Phase 1. The service structure supports future modules without implementing them now."
STATUS = "Verified: 35 backend checks, frontend build/type checks, desktop/mobile workflow tests, Docker startup, database migration and an object-storage upload/download check."
ARCH = "Browser → Next.js interface → FastAPI services → PostgreSQL"
ARCH_NOTE = "FastAPI separates API handling, business rules and database access. Certificate services use private object storage; Redis supports authentication rate limits. Each frontend page has a separate route file."
STACK = [
    ("Web interface", "Next.js, React, TypeScript, Tailwind CSS", "Pages, forms and responsive dashboards"),
    ("Backend", "Python, FastAPI", "Business rules and REST APIs"),
    ("Data tools", "SQLAlchemy, Pydantic, Alembic", "Database access, validation and schema changes"),
    ("Database", "PostgreSQL", "Employee, learning and audit records"),
    ("File storage", "S3-compatible storage; MinIO locally", "Private certificate files"),
    ("Security / rate limits", "Argon2id, JWT, refresh tokens; Redis", "Passwords, sessions and login throttling"),
    ("Local environment", "Docker Compose", "Runs the application and supporting services"),
]
DATABASE = (
    "The relational database separates login accounts from employee profiles. Core records include teams, courses, "
    "tracks, assignments, individual course attempts, certificate metadata and audit logs. Track assignments capture "
    "their course list; overlapping assignments share an unfinished course attempt, avoiding duplicate progress and "
    "evidence. Later repeat learning creates a new attempt. Certificate files stay outside PostgreSQL; only metadata "
    "and storage references are saved."
)
APIS = [
    ("/auth", "Login, session renewal, logout and current user"),
    ("/employees, /teams", "Employee profiles, history and team creation"),
    ("/courses, /learning-tracks", "Maintain courses and ordered course groups"),
    ("/assignments, /my-learning", "Assign learning and retrieve personal plans"),
    ("/completions", "Update course progress and completion"),
    ("/certificates", "Evidence upload, listing, download and revocation"),
    ("/dashboard", "Personal/team metrics and capability signals"),
    ("/audit-logs", "Review recorded administrative changes"),
]
API_NOTE = "Base path: /api/v1. APIs use GET, POST, PATCH and PUT with validation, filters, pagination and consistent errors. FastAPI enforces manager permissions and employee ownership."
SECURITY = "Passwords are hashed; sessions use HttpOnly cookies and CSRF protection. Overdue status is calculated from dates. Historical records are preserved through deactivation/cancellation. Production requires HTTPS and environment-based secrets."


def htable(headers, rows):
    head = "".join(f"<th>{escape(h)}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{escape(v)}</td>" for v in row) + "</tr>" for row in rows)
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def mtable(headers, rows):
    return "\n".join(["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"] + ["| " + " | ".join(row) + " |" for row in rows])


def paragraph(doc, text, style=None):
    p = doc.add_paragraph(text, style)
    p.paragraph_format.space_after = Pt(5)
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, (label, width) in enumerate(zip(headers, widths)):
        t.columns[i].width = Inches(width)
        t.rows[0].cells[i].text = label
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), PURPLE)
        t.rows[0].cells[i]._tc.get_or_add_tcPr().append(shade)
        for run in t.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].width = Inches(widths[i])
            if n % 2 == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F3F1EF")
                cells[i]._tc.get_or_add_tcPr().append(shade)
        no_split = OxmlElement("w:cantSplit")
        t.rows[-1]._tr.get_or_add_trPr().append(no_split)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    OUT.mkdir(exist_ok=True)
    page1 = f"""<h1>Learning Space</h1><div class="subtitle">Phase 1 · Learning Tracker | Management overview</div>
    <h2>Introduction and purpose</h2><p>{escape(INTRO)}</p>
    <h2>Users and access</h2>{htable(['Role', 'Responsibilities'], ROLES)}
    <h2>Features and modules</h2>{htable(['Module', 'What it provides'], FEATURES)}
    <p class="note">{escape(TEAM_NOTE)}</p>
    <h2>Typical workflow</h2><p>{escape(WORKFLOW)}</p>
    <h2>Business value and scope</h2><p>{escape(VALUE)}</p><p>{escape(SCOPE)}</p>
    <p class="note">{escape(STATUS)}</p>"""
    page2 = f"""<h1>Solution overview</h1><div class="subtitle">Architecture, technology, database and APIs</div>
    <h2>Architecture</h2><div class="flow">{escape(ARCH)}</div><p>{escape(ARCH_NOTE)}</p>
    <h2>Technology stack</h2>{htable(['Area', 'Technology', 'Purpose'], STACK)}
    <h2>Database design</h2><p>{escape(DATABASE)}</p>
    <h2>API overview</h2><p>{escape(API_NOTE)}</p>{htable(['API group', 'Purpose'], APIS)}
    <h2>Security and record integrity</h2><p>{escape(SECURITY)}</p>"""
    html = """<!doctype html><html lang="en"><head><meta charset="UTF-8"><title>Learning Space - Management Overview</title><style>
    @page { size: A4; margin: 0; }
    * { box-sizing: border-box; } body { margin: 0; color: #252129; font-family: Arial, sans-serif; background: #DEDAD6; }
    .sheet { width: 210mm; height: 297mm; padding: 14mm 16mm 16mm; position: relative; background: white; break-after: page; }
    .sheet:last-child { break-after: auto; } h1 { font-size: 25px; color: #4F2D7F; margin: 0 0 4px; }
    .subtitle { font-size: 11px; color: #625b68; margin-bottom: 14px; border-bottom: 3px solid #4F2D7F; padding-bottom: 9px; }
    h2 { font-size: 13.5px; color: #4F2D7F; margin: 11px 0 5px; }
    p { font-size: 12px; line-height: 1.4; margin: 4px 0 7px; }
    table { border-collapse: collapse; width: 100%; font-size: 11.5px; line-height: 1.3; margin: 5px 0 7px; }
    th { background: #4F2D7F; color: white; text-align: left; font-weight: 600; padding: 5px 7px; }
    td { padding: 5px 7px; vertical-align: top; border-bottom: 1px solid #DEDAD6; }
    tbody tr:nth-child(odd) { background: #F3F1EF; } td:first-child { width: 27%; font-weight: 600; }
    .note { font-size: 10.5px; color: #625b68; }
    .flow { background: #F3F1EF; border-left: 4px solid #4F2D7F; padding: 9px; font-size: 12px; font-weight: 600; }
    footer { position: absolute; bottom: 9mm; left: 16mm; right: 16mm; border-top: 1px solid #DEDAD6; padding-top: 5px; display: flex; justify-content: space-between; font-size: 9px; color: #625b68; }
    @media screen { .sheet { margin: 18px auto; box-shadow: 0 2px 12px #0002; } }
    </style></head><body>"""
    for n, content in enumerate((page1, page2), 1):
        html += f'<section class="sheet">{content}<footer><span>Learning Space | Phase 1 | 9 October 2026</span><span>{n} / 2</span></footer></section>'
    (OUT / "management-overview.html").write_text(html + "</body></html>", encoding="utf-8")

    md = f"# Learning Space — Phase 1 Learning Tracker\n\nManagement overview | 9 October 2026\n\n## Introduction\n\n{INTRO}\n\n## Users and access\n\n{mtable(['Role', 'Responsibilities'], ROLES)}\n\n## Features and modules\n\n{mtable(['Module', 'What it provides'], FEATURES)}\n\n{TEAM_NOTE}\n\n## Typical workflow\n\n{WORKFLOW}\n\n## Business value and scope\n\n{VALUE}\n\n{SCOPE}\n\n{STATUS}\n\n## Architecture\n\n{ARCH}\n\n{ARCH_NOTE}\n\n## Technology stack\n\n{mtable(['Area', 'Technology', 'Purpose'], STACK)}\n\n## Database design\n\n{DATABASE}\n\n## API overview\n\n{API_NOTE}\n\n{mtable(['API group', 'Purpose'], APIS)}\n\n## Security and record integrity\n\n{SECURITY}\n"
    (OUT / "management-overview.md").write_text(md, encoding="utf-8")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.55)
    section.left_margin = section.right_margin = Inches(0.63)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(9.5)
    normal.paragraph_format.line_spacing = 1.05
    for name, size in (("Title", 25), ("Heading 1", 11)):
        doc.styles[name].font.color.rgb = RGBColor.from_string(PURPLE)
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].paragraph_format.space_before = Pt(7)
        doc.styles[name].paragraph_format.space_after = Pt(4)
    doc.core_properties.title = "Learning Space - Management Overview"
    doc.core_properties.subject = "Phase 1 Learning Tracker"
    doc.core_properties.author = "Learning Space Project"
    paragraph(doc, "Learning Space", "Title")
    paragraph(doc, "Phase 1 · Learning Tracker | Management overview | 9 October 2026")
    for title, text in (("Introduction and purpose", INTRO),):
        paragraph(doc, title, "Heading 1")
        paragraph(doc, text)
    paragraph(doc, "Users and access", "Heading 1")
    table(doc, ("Role", "Responsibilities"), ROLES, (1.75, 5.25))
    paragraph(doc, "Features and modules", "Heading 1")
    table(doc, ("Module", "What it provides"), FEATURES, (1.75, 5.25))
    paragraph(doc, TEAM_NOTE)
    paragraph(doc, "Typical workflow", "Heading 1")
    paragraph(doc, WORKFLOW)
    paragraph(doc, "Business value and scope", "Heading 1")
    paragraph(doc, VALUE)
    paragraph(doc, SCOPE)
    paragraph(doc, STATUS)
    doc.add_page_break()
    paragraph(doc, "Solution overview", "Title")
    paragraph(doc, "Architecture, technology, database and APIs")
    paragraph(doc, "Architecture", "Heading 1")
    paragraph(doc, ARCH)
    paragraph(doc, ARCH_NOTE)
    paragraph(doc, "Technology stack", "Heading 1")
    table(doc, ("Area", "Technology", "Purpose"), STACK, (1.35, 3.35, 2.3))
    paragraph(doc, "Database design", "Heading 1")
    paragraph(doc, DATABASE)
    paragraph(doc, "API overview", "Heading 1")
    paragraph(doc, API_NOTE)
    table(doc, ("API group", "Purpose"), APIS, (2.3, 4.7))
    paragraph(doc, "Security and record integrity", "Heading 1")
    paragraph(doc, SECURITY)
    foot = section.footer.paragraphs[0]
    foot.text = "Learning Space | Phase 1 | Management overview"
    foot.runs[0].font.size = Pt(8)
    foot.runs[0].font.color.rgb = RGBColor.from_string(PURPLE)
    doc.save(OUT / "management-overview.docx")
    print("Created management-overview.md, .html and .docx")


if __name__ == "__main__":
    main()
